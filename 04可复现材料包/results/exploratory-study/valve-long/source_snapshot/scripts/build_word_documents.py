"""Build the two submission Word documents from their Markdown sources.

Use the Codex bundled Python runtime; provide a portable Pandoc executable with
--pandoc. Equations are converted by Pandoc to editable Office Math (OMML).
The source data are read only. Figure variants use Times New Roman labels.
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
SOURCES = {
    "coupling-design.md": "一三维耦合方案说明.docx",
    "case-validation-report.md": "一三维双向耦合算例验证报告.docx",
}
MATH_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
WORD_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
VARIABLES = re.compile(r"(?<![A-Za-z0-9_])(alpha|H|T|R|x|y|z|L1|p50|p95|p99)(?![A-Za-z0-9_])")
SYMBOLS = {"alpha": r"\alpha", "L1": "L^1", "p50": "p_{50}",
           "p95": "p_{95}", "p99": "p_{99}"}


def read_csv(path):
    with path.open(encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def make_figures():
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    import numpy as np
    from matplotlib import font_manager
    font_manager.fontManager.addfont(r"C:\Windows\Fonts\times.ttf")
    plt.rcParams.update({"font.family": "Times New Roman", "font.size": 13,
                         "axes.labelsize": 14, "legend.fontsize": 11,
                         "axes.unicode_minus": False})
    source = ROOT / "results/v03-study"
    assets = DOCS / "assets/word-report"
    assets.mkdir(parents=True, exist_ok=True)
    result = {}

    # Architecture uses words rather than mathematical symbols in the bitmap.
    from matplotlib.patches import FancyBboxPatch, FancyArrowPatch
    cn = font_manager.FontProperties(fname=r"C:\Windows\Fonts\simsun.ttc", size=15)
    fig, ax = plt.subplots(figsize=(9.4, 3.6))
    ax.set_xlim(0, 10); ax.set_ylim(0, 4); ax.axis("off")
    for x, en, zh, description in [(.4, "Python", "一维水力与混合状态", "管段惯性与混合容积"),
                                   (3.65, "preCICE", "耦合协调与收敛判定", "窗口迭代与检查点回滚"),
                                   (6.9, "OpenFOAM", "三维流场与标量输运", "速度压力与实际面通量")]:
        ax.add_patch(FancyBboxPatch((x, 1.1), 2.7, 1.65, boxstyle="round,pad=0.06",
                                   facecolor="#eef2f6", edgecolor="#687d8f", linewidth=1.1))
        ax.text(x+1.35, 2.45, en, ha="center", va="center", fontsize=16)
        ax.text(x+1.35, 1.96, zh, ha="center", va="center", fontproperties=cn)
        ax.text(x+1.35, 1.47, description, ha="center", va="center", fontproperties=cn, fontsize=12)
    for start, end in [((3.16, 2.25), (3.56, 2.25)), ((6.41, 2.25), (6.8, 2.25))]:
        ax.add_patch(FancyArrowPatch(start, end, arrowstyle="-|>", mutation_scale=13, color="#3e668a"))
    ax.text(5, 3.18, "正向交换平均速度 温度与相分数", ha="center", fontproperties=cn)
    ax.annotate("", xy=(1.75, .75), xytext=(8.25, .75),
                arrowprops={"arrowstyle": "-|>", "color": "#a55e37", "lw": 1.5})
    ax.plot([1.75, 1.75], [.75, 1.04], color="#a55e37")
    ax.plot([8.25, 8.25], [.75, 1.04], color="#a55e37")
    ax.text(5, .3, "三维接口背压反馈至一维模型", ha="center", fontproperties=cn)
    fig.tight_layout(pad=.3)
    result["assets/word-report/coupling-workflow.png"] = assets / "coupling-workflow.png"
    fig.savefig(assets / "coupling-workflow.png", dpi=220); plt.close(fig)

    fig, axes = plt.subplots(1, 2, figsize=(9.4, 3.7), constrained_layout=True)
    labels = ["8 / 40", "12 / 80"]
    x = np.arange(2); width = .34
    old, new, pressure = [], [], []
    for name in ["steady", "steady-fine"]:
        former = json.loads((ROOT / "results/v02-study" / name / "summary.json").read_text())
        current = json.loads((source / name / "summary.json").read_text())
        old.append(100*former["steady_poiseuille_reference"]["relative_final_velocity_error"])
        new.append(100*current["steady_poiseuille_reference"]["relative_final_velocity_error"])
        pressure.append(100*current["steady_poiseuille_reference"]["relative_final_pressure_error"])
    for ax, left, right, legend in [(axes[0], old, new, ("v0.2 velocity", "v0.3 velocity")),
                                    (axes[1], new, pressure, ("v0.3 velocity", "v0.3 pressure"))]:
        ax.bar(x-width/2, left, width, label=legend[0], color="#476e95")
        ax.bar(x+width/2, right, width, label=legend[1], color="#bf7d46")
        ax.set_xticks(x, labels); ax.set_xlabel("Radial / axial grid cells")
        ax.set_ylabel("Relative error [%]"); ax.grid(axis="y", alpha=.25); ax.legend(fontsize=11)
    result["assets/word-report/steady-error-comparison.png"] = assets / "steady-error-comparison.png"
    fig.savefig(assets / "steady-error-comparison.png", dpi=220); plt.close(fig)

    balance = json.loads((source / "valve-0.001/transport_summary.json").read_text())["interface_window_integrals"]
    keys = ["mass", "thermal", "phase", "pressure_force", "pressure_work"]
    names = ["Mass", "Sensible heat", "Phase A", "Pressure impulse", "Pressure work"]
    y = np.arange(len(keys)); height = .33
    cumulative = [100*balance[k]["relative_cumulative_residual"] for k in keys]
    absolute = [100*balance[k]["relative_l1_window_residual"] for k in keys]
    fig, ax = plt.subplots(figsize=(8.7, 3.8), constrained_layout=True)
    ax.barh(y-height/2, cumulative, height, label="Cumulative residual", color="#476e95")
    ax.barh(y+height/2, absolute, height, label="Absolute window sum", color="#bf7d46")
    ax.set_yticks(y, names); ax.invert_yaxis()
    ax.set_xlabel("Relative interface residual [%]"); ax.grid(axis="x", alpha=.25)
    ax.axvline(1, color="gray", linestyle="--", linewidth=.9, label="1 percent reference")
    ax.legend(fontsize=11, loc="lower right", bbox_to_anchor=(1, 1.01), ncol=2)
    result["assets/word-report/window-conservation.png"] = assets / "window-conservation.png"
    fig.savefig(assets / "window-conservation.png", dpi=220); plt.close(fig)

    # Same numerical samples and reference as the source pressure figure.
    base = source / "steady-fine"
    case = json.loads((base / "case.json").read_text(encoding="utf-8"))
    rho = case["pipe"]["rho"]
    samples = list((base / "fluid3d/postProcessing/sampleDict").glob("*/centerline_p.xy"))
    cfd = np.loadtxt(max(samples, key=lambda p: float(p.parent.name)), ndmin=2)
    one = np.loadtxt(base / "pipe_final.csv", delimiter=",", skiprows=1, ndmin=2)
    length = case["pipe"]["length"] + case["cfd"]["length"]
    fig, ax = plt.subplots(figsize=(8.7, 3.7), constrained_layout=True)
    ax.plot(one[:, 0], one[:, 1], label="1D")
    ax.plot(cfd[:, 0], cfd[:, 1] * rho, label="3D centreline")
    x = np.linspace(0, length, 101)
    ax.plot(x, case["pipe"]["supply_pressure"] * (1-x/length), "k--", label="Steady Poiseuille")
    ax.axvline(case["pipe"]["length"], color="gray", linestyle=":", label="Coupling interface")
    ax.set_xlabel("Axial coordinate [m]"); ax.set_ylabel("Gauge pressure [Pa]")
    ax.grid(alpha=.25); ax.legend()
    result["../results/v03-study/steady-fine/pressure-profile.png"] = assets / "pressure-profile.png"
    fig.savefig(assets / "pressure-profile.png", dpi=220); plt.close(fig)

    fig, axes = plt.subplots(1, 2, figsize=(9.4, 3.7), constrained_layout=True)
    for name, label in [("advection", "6 x 40, time step 0.5 ms"),
                        ("advection-fine", "8 x 80, time step 0.25 ms")]:
        rows = read_csv(source / name / "foam_accepted.csv")
        t = [float(r["time_s"]) for r in rows]
        axes[0].plot(t, [float(r["temperature_rms_k"]) for r in rows], label=label)
        axes[1].plot(t, [float(r["alpha_rms"]) for r in rows], label=label)
    axes[0].set_ylabel("Temperature field RMS [K]")
    axes[1].set_ylabel("Phase A field RMS")
    for ax in axes:
        ax.set_xlabel("Physical time [s]"); ax.grid(alpha=.25); ax.legend(fontsize=11)
    result["../results/v03-study/advection-refinement.png"] = assets / "advection-refinement.png"
    fig.savefig(assets / "advection-refinement.png", dpi=220); plt.close(fig)

    frequency = read_csv(source / "frequency.csv")
    fig, axes = plt.subplots(1, 2, figsize=(9.4, 3.7), constrained_layout=True)
    x = [float(r["frequency_hz"]) for r in frequency]
    for field, label in [("velocity_m_s", "Velocity"), ("back_pressure_pa", "Back pressure"),
                         ("temperature_k", "Temperature"), ("phase_fraction", "Phase fraction")]:
        axes[0].plot(x, [100*float(r["normalized_rms_"+field]) for r in frequency], "o-", label=label)
    axes[0].set_ylabel("RMS / reference signal range [%]"); axes[0].legend(fontsize=11)
    axes[1].plot(x, [float(r["wall_seconds"]) for r in frequency], "o-")
    axes[1].set_ylabel("Wall time [s], one run per point")
    for ax in axes:
        ax.set_xscale("log"); ax.set_xlabel("Exchange frequency [Hz]"); ax.grid(alpha=.25)
    result["../results/v03-study/frequency.png"] = assets / "frequency.png"
    fig.savefig(assets / "frequency.png", dpi=220); plt.close(fig)

    fig, axes = plt.subplots(1, 2, figsize=(9.4, 3.7), constrained_layout=True)
    for record in frequency:
        rows = read_csv(source / record["name"] / "pipe.csv")
        t = np.array([float(r["time_s"]) for r in rows])
        pressure = np.array([float(r["back_pressure_pa"]) for r in rows])
        label = f"Window {float(record['window_s'])*1000:g} ms"
        axes[0].plot(t*1000, pressure, label=label)
        axes[1].plot(t*1000, pressure, ".-", markersize=3, label=label)
    axes[1].set_xlim(49.8, 50.6)
    for ax in axes:
        ax.set_xlabel("Physical time [ms]"); ax.set_ylabel("Back pressure [Pa]"); ax.grid(alpha=.25)
    axes[0].legend(fontsize=11)
    result["../results/v03-study/pressure-event-comparison.png"] = assets / "pressure-event-comparison.png"
    fig.savefig(assets / "pressure-event-comparison.png", dpi=220); plt.close(fig)
    return result


def str_inlines(text):
    result = []
    cursor = 0
    for match in VARIABLES.finditer(text):
        if match.start() > cursor:
            result.append({"t": "Str", "c": text[cursor:match.start()]})
        result.append({"t": "Math", "c": [{"t": "InlineMath"}, SYMBOLS.get(match[0], match[0])]})
        cursor = match.end()
    if cursor < len(text):
        result.append({"t": "Str", "c": text[cursor:]})
    return result or [{"t": "Str", "c": text}]


def rewrite_ast(node, figures):
    if isinstance(node, list):
        result = []
        for child in node:
            if isinstance(child, dict) and child.get("t") == "Str":
                result.extend(str_inlines(child["c"]))
            else:
                result.append(rewrite_ast(child, figures))
        return result
    if not isinstance(node, dict):
        return node
    kind = node.get("t")
    # Software names, code, path targets and existing math must stay intact.
    if kind in {"Code", "Math"}:
        return node
    if kind == "Header":
        node["c"][0] = max(1, node["c"][0]-1)
    if kind == "Link":
        node["c"][1] = rewrite_ast(node["c"][1], figures)
        target = node["c"][2][0]
        if target in SOURCES:
            node["c"][2][0] = SOURCES[target]
        return node
    if kind == "Image":
        target = node["c"][2][0]
        if target in figures:
            node["c"][2][0] = str(figures[target])
        return node
    return {key: rewrite_ast(value, figures) for key, value in node.items()}


def architecture_blocks():
    steps = [
        "控制时钟及事件表驱动 Python 一维惯性管段与混合容积。",
        "一维通过 preCICE 写出 Velocity、FlowTemperature 和 Alpha，三维 OpenFOAM 应用入口边界并推进速度、压力及标量方程。",
        "三维通过 Pressure 返回接口背压，一维读入反馈；未收敛时恢复检查点并重复试算。",
        "收敛后提交一维状态及账本，与三维面通量和接受轨迹共同用于精度、守恒、延迟与效率分析。",
    ]
    return [{"t": "OrderedList", "c": [[1, {"t": "Decimal"}, {"t": "Period"}],
              [[{"t": "Plain", "c": [{"t": "Str", "c": step}]}] for step in steps]]}]


def set_fonts(rpr):
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts"); rpr.insert(0, fonts)
    for attr in list(fonts.attrib):
        del fonts.attrib[attr]
    for name in ("ascii", "hAnsi", "cs"):
        fonts.set(qn("w:"+name), "Times New Roman")
    fonts.set(qn("w:eastAsia"), "宋体")
    color = rpr.find(qn("w:color"))
    if color is None:
        color = OxmlElement("w:color"); rpr.append(color)
    color.attrib.clear(); color.set(qn("w:val"), "000000")


def set_run(run, size=None):
    set_fonts(run._element.get_or_add_rPr())
    if size:
        run.font.size = Pt(size)


def style_document(path, title):
    doc = Document(path)
    section = doc.sections[0]
    section.page_width = Cm(21); section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0); section.bottom_margin = Cm(1.9)
    section.left_margin = Cm(2.2); section.right_margin = Cm(2.2)
    section.footer_distance = Cm(.85)
    for style in doc.styles:
        set_fonts(style.element.get_or_add_rPr())
        if style.type == 1:
            style.paragraph_format.widow_control = True
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.keep_with_next = False
    for name in ("Body Text", "First Paragraph", "Compact"):
        if name in doc.styles:
            style = doc.styles[name]
            style.base_style = normal
            style.font.size = Pt(11)
            style.paragraph_format.line_spacing = 1.25
            style.paragraph_format.space_after = Pt(5)
    for name, size, before, after in [("Title", 20, 0, 10), ("Heading 1", 14, 12, 6),
                                      ("Heading 2", 12, 9, 5), ("Heading 3", 11, 7, 4)]:
        style = doc.styles[name]
        style.font.size = Pt(size); style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True
        style.paragraph_format.line_spacing = 1.15
    doc.paragraphs[0].style = doc.styles["Title"]
    doc.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for p in doc.paragraphs:
        if p.style.name not in {"Title", "Heading 1", "Heading 2", "Heading 3", "Image Caption"}:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for run in p.runs:
            set_run(run)
        # Long identifiers wrap visibly without changing the identifier text.
        if "sha256:" in p.text:
            wrap = OxmlElement("w:wordWrap"); wrap.set(qn("w:val"), "1")
            p._p.get_or_add_pPr().append(wrap)
            for run in p.runs:
                run.font.size = Pt(9)
        if p._p.xpath(".//w:drawing"):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = Pt(7)
            p.paragraph_format.space_after = Pt(4)
        if p.style.name == "Image Caption":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.keep_with_next = False
            for run in p.runs:
                set_run(run, 10)
        if p._p.xpath(".//m:oMathPara"):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing = 1.05
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_together = True
    # Constrain drawings to the printable width; preserve their aspect ratio.
    for shape in doc.inline_shapes:
        if shape.width > Cm(16.4):
            ratio = Cm(16.4) / shape.width
            shape.width = int(shape.width*ratio); shape.height = int(shape.height*ratio)

    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        count = len(table.columns)
        if count == 2:
            widths = [4.2, 12.4]
        elif count == 3:
            # Conservation formulas need a narrow label column and wide formula columns.
            if "一维接口积分" in table.rows[0].cells[1].text:
                widths = [3.1, 6.75, 6.75]
            elif table.rows[0].cells[0].text == "指标":
                widths = [3.1, 6.6, 6.9]
            else:
                widths = [6.8, 4.9, 4.9]
        elif count == 4:
            widths = [3.5, 2.0, 5.1, 6.0] if table.rows[0].cells[0].text == "字段" else [4.0, 3.8, 3.8, 5.0]
        elif count == 5:
            widths = [4.0, 3.15, 3.15, 3.15, 3.15]
        elif count == 6:
            widths = [3.2, 1.8, 2.9, 2.9, 2.9, 2.9]
        elif count == 7:
            widths = [2.0, 1.9, 2.5, 2.5, 2.5, 2.6, 2.6]
        else:
            widths = [16.6/count]*count
        for column, width in zip(table.columns, widths):
            column.width = Cm(width)
        pr = table._tbl.tblPr
        borders = pr.find(qn("w:tblBorders"))
        if borders is None:
            borders = OxmlElement("w:tblBorders"); pr.append(borders)
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement("w:"+edge)
            for attr, value in {"val": "single", "sz": "4", "color": "D9D9D9"}.items():
                el.set(qn("w:"+attr), value)
            borders.append(el)
        margins = OxmlElement("w:tblCellMar")
        for side, value in [("top", 75), ("bottom", 75), ("left", 85), ("right", 85)]:
            el = OxmlElement("w:"+side); el.set(qn("w:w"), str(value)); el.set(qn("w:type"), "dxa")
            margins.append(el)
        pr.append(margins)
        for i, row in enumerate(table.rows):
            trpr = row._tr.get_or_add_trPr()
            if i == 0:
                trpr.append(OxmlElement("w:tblHeader"))
            trpr.append(OxmlElement("w:cantSplit"))
            for j, cell in enumerate(row.cells):
                cell.width = Cm(widths[j])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "E7ECF0" if i == 0 else "FFFFFF")
                cell._tc.get_or_add_tcPr().append(shade)
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.line_spacing = 1.15
                    p.paragraph_format.keep_with_next = i == 0
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 or count <= 3 else WD_ALIGN_PARAGRAPH.CENTER
                    for run in p.runs:
                        set_run(run, 9.5 if count >= 5 else 10)
                        run.font.bold = i == 0
    # Apply explicit script font mappings to ALL runs including hyperlinks and OMML.
    for run in doc.element.xpath(".//w:r"):
        rpr = run.find(qn("w:rPr"))
        if rpr is None:
            rpr = OxmlElement("w:rPr"); run.insert(0, rpr)
        set_fonts(rpr)
    for run in doc.element.xpath(".//m:r"):
        rpr = run.find(qn("w:rPr"))
        if rpr is None:
            rpr = OxmlElement("w:rPr"); run.insert(0, rpr)
        set_fonts(rpr)
        sz = OxmlElement("w:sz"); sz.set(qn("w:val"), "22"); rpr.append(sz)
        # Inline equations in tables inherit the table's compact point size.
        if any(el.tag == qn("w:tc") for el in run.iterancestors()):
            sz.set(qn("w:val"), "19")
    settings = doc.settings.element
    mathpr = settings.find(qn("m:mathPr"))
    if mathpr is None:
        mathpr = OxmlElement("m:mathPr"); settings.append(mathpr)
    font = mathpr.find(qn("m:mathFont"))
    if font is None:
        font = OxmlElement("m:mathFont"); mathpr.insert(0, font)
    font.set(qn("m:val"), "Times New Roman")
    # Remove direct title border residue and template theme fonts everywhere.
    for border in doc.paragraphs[0]._p.xpath("./w:pPr/w:pBdr"):
        border.getparent().remove(border)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(footer.add_run("第 "), 9)
    fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), "PAGE")
    footer._p.append(fld)
    set_run(footer.add_run(" 页"), 9)
    doc.core_properties.title = title
    doc.core_properties.subject = "一维系统与三维 CFD 双向耦合"
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.keywords = "OpenFOAM preCICE v0.3"
    doc.save(path)


def build(pandoc, figures, exploration_root=None):
    with tempfile.TemporaryDirectory(prefix="cfd_docx_build_") as tmp:
        for source_name, output_name in SOURCES.items():
            text = (DOCS / source_name).read_text(encoding="utf-8")
            if exploration_root:
                supplement = "scheme-section.md" if source_name == "coupling-design.md" else "report-section.md"
                text += "\n" + (exploration_root / supplement).read_text(encoding="utf-8")
                text = text.replace("版本：v0.3；日期：2026-10-10。", "版本：v0.3 探索增强版；日期：2026-10-11。", 1)
            title = text.splitlines()[0].lstrip("# ")
            # Vertical bars inside formulas must not split Markdown table cells.
            text = re.sub(r"\$\$(.*?)\$\$|\$([^$\n]+)\$", lambda m:
                          ("$$"+m[1].replace("|", r"\vert{}")+"$$") if m[1] is not None
                          else ("$"+m[2].replace("|", r"\vert{}")+"$"), text, flags=re.S)
            # Two formulas in one display are easier to read on separate lines.
            text = re.sub(r"\$\$(.*?)\$\$", lambda m: "\n\n".join(
                "$$"+part.strip().rstrip(",")+"$$" for part in re.split(r"\\qquad\s*", m[1])), text, flags=re.S)
            ast = json.loads(subprocess.run([pandoc, "-f", "markdown", "-t", "json"],
                                            input=text, text=True, encoding="utf-8", capture_output=True, check=True).stdout)
            blocks = []
            for block in ast["blocks"]:
                if block["t"] == "CodeBlock" and "mermaid" in block["c"][0][1]:
                    blocks.extend(architecture_blocks())
                else:
                    blocks.append(block)
            ast["blocks"] = blocks
            ast = rewrite_ast(ast, figures)
            temp_docx = Path(tmp) / output_name
            subprocess.run([pandoc, "-f", "json", "-t", "docx", "--standalone",
                            "--resource-path", str(DOCS), "-o", str(temp_docx)],
                           input=json.dumps(ast, ensure_ascii=False), text=True, encoding="utf-8", check=True)
            style_document(temp_docx, title)
            target = DOCS / output_name
            target.write_bytes(temp_docx.read_bytes())
            print(target)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--pandoc", required=True, type=Path)
    parser.add_argument("--python-lib-dir", type=Path, help="Optional temporary figure dependencies")
    parser.add_argument("--exploration-root", type=Path, help="Verified exploration sections; creates enhanced copies")
    args = parser.parse_args()
    if args.python_lib_dir:
        import sys
        sys.path.insert(0, str(args.python_lib_dir))
    if args.exploration_root:
        SOURCES.update({key: Path(value).stem + "_探索增强版.docx" for key, value in SOURCES.items()})
    build(str(args.pandoc.resolve()), make_figures(), args.exploration_root.resolve() if args.exploration_root else None)
